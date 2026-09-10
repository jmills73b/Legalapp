"""DOCX output on the firm's house style.

Two modes, both useful:

- unmerged: placeholders are kept and highlighted, so a fee earner can fill the
  letter in in Word and see at a glance what is left;
- merged: values supplied, nothing highlighted, ready to sign.

House style lives in library/house_style.yaml so a firm can change how its
letters look without touching this file.
"""
from __future__ import annotations

import pathlib
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Cm, Pt

from . import library
from .render import merge
from .tokens import TOKEN_RE

_SUBJECT_LINE = re.compile(r"^Re:\s", re.I)
_LABEL_LINE = re.compile(r"^([A-Z][A-Za-z ]{0,20}):\s+(.*)$")
_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


def style() -> dict:
    return library._load("house_style.yaml")


def _setup(doc: Document, st: dict) -> None:
    section = doc.sections[0]
    page = st["page"]
    if page.get("size", "A4").upper() == "A4":
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = Cm(page["margin_top_cm"])
    section.bottom_margin = Cm(page["margin_bottom_cm"])
    section.left_margin = Cm(page["margin_left_cm"])
    section.right_margin = Cm(page["margin_right_cm"])

    normal = doc.styles["Normal"]
    normal.font.name = st["body"]["font"]
    normal.font.size = Pt(st["body"]["size_pt"])
    pf = normal.paragraph_format
    pf.line_spacing = st["body"]["line_spacing"]
    pf.space_after = Pt(st["body"]["space_after_pt"])
    pf.alignment = _ALIGN.get(st["body"]["alignment"], WD_ALIGN_PARAGRAPH.LEFT)


def _add_runs(paragraph, text: str, highlight: bool) -> None:
    """Write text, highlighting any placeholder still in it."""
    if not highlight or not TOKEN_RE.search(text):
        paragraph.add_run(text)
        return
    pos = 0
    for m in TOKEN_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        run = paragraph.add_run(m.group(0))
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _ref_line(doc: Document, line: str, st: dict, highlight: bool) -> None:
    """"Our ref:" and its value separated by a tab stop, not by spaces."""
    cfg = st["ref_block"]
    m = _LABEL_LINE.match(line)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(cfg["tab_stop_cm"]))
    _add_runs(p, f"{m.group(1)}:\t{m.group(2)}" if m else line, highlight)
    for run in p.runs:
        run.font.size = Pt(cfg["size_pt"])


def _heading(doc: Document, text: str, st: dict, highlight: bool,
             letterhead: bool = False) -> None:
    """The firm name, the ref block and the Re: line all arrive as headings.

    Lines are dispatched in source order -- grouping them would reorder the
    letterhead.
    """
    lines = [ln for ln in text.splitlines() if ln.strip()]
    for i, line in enumerate(lines):
        if _SUBJECT_LINE.match(line):
            p = doc.add_paragraph()
            _add_runs(p, line, highlight)
            for run in p.runs:
                run.bold = st["subject"]["bold"]
            p.paragraph_format.space_before = Pt(st["subject"]["space_before_pt"])
            p.paragraph_format.space_after = Pt(st["subject"]["space_after_pt"])
        elif _LABEL_LINE.match(line):
            _ref_line(doc, line, st, highlight)
        elif letterhead and i == 0:
            cfg = st["firm_name"]
            p = doc.add_paragraph()
            _add_runs(p, line, highlight)
            for run in p.runs:
                run.bold = cfg["bold"]
                run.font.size = Pt(cfg["size_pt"])
            p.paragraph_format.space_after = Pt(cfg["space_after_pt"])
        else:
            p = doc.add_paragraph()
            _add_runs(p, line, highlight)
            p.paragraph_format.space_after = Pt(2 if letterhead else st["body"]["space_after_pt"])


def _closing(doc: Document, text: str, st: dict, highlight: bool) -> None:
    """Sign-off, a gap to sign in, then the name block."""
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines:
        return
    p = doc.add_paragraph()
    _add_runs(p, lines[0], highlight)
    p.paragraph_format.space_after = Pt(0)
    for _ in range(st["signature_gap_lines"]):
        gap = doc.add_paragraph()
        gap.paragraph_format.space_after = Pt(0)
    for line in lines[1:]:
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(0)
        _add_runs(q, line, highlight)


def write(letter, path: str | pathlib.Path, values: dict[str, str] | None = None,
          letter_date=None) -> pathlib.Path:
    """Render a composed Letter to .docx. With values, it is merged and ready
    to sign; without, placeholders are kept and highlighted."""
    st = style()
    doc = Document()
    _setup(doc, st)
    highlight = st.get("highlight_unfilled_tokens", True) and not values

    seen_heading = False
    for para in letter.paragraphs:
        text = para.text.strip()
        if values is not None:
            text, _ = merge(text, values, letter.spec, letter_date)
            text = text.strip()
        if not text:
            continue

        if para.role == "heading":
            _heading(doc, text, st, highlight, letterhead=not seen_heading)
            seen_heading = True
        elif para.role == "marking":
            p = doc.add_paragraph()
            _add_runs(p, text, highlight)
            for run in p.runs:
                run.bold = st["marking"]["bold"]
            p.paragraph_format.space_after = Pt(st["marking"]["space_after_pt"])
        elif para.role == "closing":
            _closing(doc, text, st, highlight)
        elif para.role == "enclosure":
            p = doc.add_paragraph()
            _add_runs(p, text, highlight)
            for run in p.runs:
                run.font.size = Pt(st["enclosure"]["size_pt"])
            p.paragraph_format.space_before = Pt(st["enclosure"]["space_before_pt"])
        else:
            p = doc.add_paragraph()
            _add_runs(p, " ".join(text.split()), highlight)

    path = pathlib.Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
