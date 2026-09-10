"""DOCX output: house style applied, placeholders visible until they are filled."""
import pytest

docx_lib = pytest.importorskip("docx")
from docx import Document

from legalapp import docx as dx
from legalapp import library
from legalapp.compose import compose
from legalapp.spec import LetterSpec, RecipientClass


def letter(lt_id="child_arrangements_proposal"):
    lt = library.letter_types()[lt_id]
    return compose(LetterSpec(lt_id, RecipientClass(lt["recipient_class"]),
                              deadline_days=lt.get("default_deadline_days"),
                              privileged=lt["wp_marking"] == "required"))


VALUES = {
    "FIRM_NAME": "Harbrook & Vane LLP", "FIRM_ADDRESS_BLOCK": "12 Bedford Row, London WC1R 4BU",
    "MATTER_REF": "HV/FAM/2261", "THEIR_REF": "KDR/4417", "RECIPIENT_SALUTATION": "Sirs",
    "CLIENT_FULL_NAME": "Ayesha Kaur", "OTHER_PARTY_NAME": "Daniel Kaur",
    "FEE_EARNER_NAME": "J. Mills", "FEE_EARNER_ROLE": "Senior Associate",
}


def _paras(path):
    return Document(str(path)).paragraphs


def test_unfilled_placeholders_are_highlighted(tmp_path):
    paras = _paras(dx.write(letter(), tmp_path / "a.docx"))
    highlighted = {r.text for p in paras for r in p.runs if r.font.highlight_color is not None}
    assert "[CLIENT_FULL_NAME]" in highlighted
    assert all(t.startswith("[") and t.endswith("]") for t in highlighted)


def test_merged_letter_highlights_nothing(tmp_path):
    paras = _paras(dx.write(letter(), tmp_path / "b.docx", values=VALUES))
    assert not [r for p in paras for r in p.runs if r.font.highlight_color is not None]
    assert any("Ayesha Kaur" in p.text for p in paras)
    assert not any("[" in p.text for p in paras)


def test_without_prejudice_marking_is_bold_and_first(tmp_path):
    paras = _paras(dx.write(letter(), tmp_path / "c.docx"))
    assert paras[0].text == "WITHOUT PREJUDICE"
    assert all(r.bold for r in paras[0].runs)


def test_subject_line_is_bold_not_tab_stopped(tmp_path):
    # "Re:" matches the ref-block label pattern; it must be classified as a
    # subject line first, or it renders tabbed and unbolded.
    paras = _paras(dx.write(letter(), tmp_path / "d.docx"))
    subject = next(p for p in paras if p.text.startswith("Re:"))
    assert any(r.bold for r in subject.runs)
    assert "\t" not in subject.text


def test_ref_block_uses_a_tab_stop(tmp_path):
    paras = _paras(dx.write(letter(), tmp_path / "e.docx"))
    ref = next(p for p in paras if p.text.startswith("Our ref:"))
    assert "\t" in ref.text
    assert list(ref.paragraph_format.tab_stops)


def test_firm_name_carries_the_letterhead_style(tmp_path):
    st = dx.style()
    paras = _paras(dx.write(letter(), tmp_path / "f.docx"))
    firm = next(p for p in paras if "[FIRM_NAME]" in p.text)
    assert any(r.font.size and r.font.size.pt == st["firm_name"]["size_pt"] for r in firm.runs)


def test_signature_gap_is_left_to_sign_in(tmp_path):
    st = dx.style()
    paras = _paras(dx.write(letter(), tmp_path / "g.docx", values=VALUES))
    idx = next(i for i, p in enumerate(paras) if p.text.strip() == "Yours faithfully")
    gap = [p for p in paras[idx + 1:idx + 1 + st["signature_gap_lines"]] if not p.text.strip()]
    assert len(gap) == st["signature_gap_lines"]
    assert "J. Mills" in paras[idx + 1 + st["signature_gap_lines"]].text


def test_page_is_a4_with_house_margins(tmp_path):
    st = dx.style()
    section = Document(str(dx.write(letter(), tmp_path / "h.docx"))).sections[0]
    assert round(section.page_width.cm, 1) == 21.0
    assert round(section.top_margin.cm, 1) == st["page"]["margin_top_cm"]


@pytest.mark.parametrize("lt_id", list(library.letter_types()))
def test_every_letter_type_renders(tmp_path, lt_id):
    paras = _paras(dx.write(letter(lt_id), tmp_path / f"{lt_id}.docx"))
    assert [p for p in paras if p.text.strip()]
